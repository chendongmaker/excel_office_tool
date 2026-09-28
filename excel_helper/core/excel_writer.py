from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_table(frame: pd.DataFrame, output_path: str | Path, sheet_name: str = "结果") -> Path:
    """把 DataFrame 导出为带基础样式的 Excel。

    上游：`ProcessService` 在合并、清洗等任务完成后调用。
    下游：返回输出路径，UI 记录到历史任务并弹窗提示用户。
    """

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        # 表头样式统一放在写入层，避免每个业务功能重复处理 Excel 美化。
        header_fill = PatternFill("solid", fgColor="EAF2FF")
        header_font = Font(bold=True, color="1F2937")

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            # 自动列宽限制在 10-36，兼顾可读性和大字段不撑爆页面。
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 10), 36)

        # 冻结首行和启用筛选，方便用户打开结果文件后继续办公。
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    return path
